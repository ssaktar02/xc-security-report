"""
Word (.docx) exporter — produces a client-facing "Best Practice Guide" that
can be used to manually verify the security posture of an F5 XC deployment.

The document has two parts:

  Part 1 — Generic Best Practice Checklist
      Every distinct security check (deduplicated by rule id) plus the manual
      review items, grouped by category. Each row carries a printable checkbox
      so the client can tick it off during their own review. This part is the
      same for every client.

  Part 2 — Per-Report Appendix
      For each assessed Load Balancer, the actual findings from the uploaded
      audit (status, object, finding, recommendation) with a verification
      checkbox — so the client can confirm remediation for their own tenant.

Design notes:
  - Underlying status values (FAIL/WARN) are shown with the customer-facing
    display labels (REVIEW/CAUTION) via report_data.status_display_label.
  - Checkboxes are the printable Unicode ballot box (U+2610); the client can
    tick them on paper or click and type inside Word.
"""

import io
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import report_data as rd

CHECKBOX = "\u2610"  # ☐ ballot box

# Brand-ish palette (kept close to the HTML/PDF report)
_NAVY = RGBColor(0x1A, 0x1A, 0x2E)
_RED = RGBColor(0xE2, 0x23, 0x1A)
_ORANGE = RGBColor(0xE6, 0x7A, 0x22)
_GREEN = RGBColor(0x27, 0xAE, 0x60)
_GREY = RGBColor(0x66, 0x66, 0x66)
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

_SEVERITY_ORDER = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2, "LOW": 3, "INFO": 4}


# ----------------------------------------------------------------------
# Low-level docx helpers
# ----------------------------------------------------------------------

def _shade_cell(cell, hex_color: str):
    """Apply a solid background fill to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _set_cell_text(cell, text, bold=False, color=None, size=9, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run(text if text is not None else "")
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return p


def _status_color(status: str, severity: str) -> RGBColor:
    s = (status or "").upper()
    if s == "FAIL":
        return _RED
    if s == "WARN":
        return _ORANGE
    if s == "PASS":
        return _GREEN
    return _GREY


def _heading(doc, text, level=1, color=_NAVY):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.font.color.rgb = color
    return h


def _set_col_widths(table, widths):
    """Force fixed column widths (python-docx needs it set per-cell)."""
    table.autofit = False
    for row in table.rows:
        for idx, w in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = w


def _header_row(table, labels):
    hdr = table.rows[0].cells
    for i, label in enumerate(labels):
        _set_cell_text(hdr[i], label, bold=True, color=_WHITE, size=9)
        _shade_cell(hdr[i], "1A1A2E")


# ----------------------------------------------------------------------
# Content builders
# ----------------------------------------------------------------------

def _category_label(cat: str) -> str:
    if not cat or cat.upper() == "UNCATEGORIZED":
        return "General"
    return cat.replace("_", " ").title()


def _dedup_rules(all_findings):
    """One entry per rule id, preferring the most severe/actionable instance."""
    best = {}
    for f in all_findings:
        rid = f.rule_id or ""
        if not rid:
            continue
        prev = best.get(rid)
        if prev is None:
            best[rid] = f
            continue
        # Prefer a non-PASS instance and higher severity
        prev_pass = prev.status.upper() == "PASS"
        cur_pass = f.status.upper() == "PASS"
        if prev_pass and not cur_pass:
            best[rid] = f
        elif prev_pass == cur_pass:
            if _SEVERITY_ORDER.get(f.severity.upper(), 9) < _SEVERITY_ORDER.get(prev.severity.upper(), 9):
                best[rid] = f
    return list(best.values())


def _add_generic_guide(doc, all_findings, manual_items):
    _heading(doc, "Part 1 — Security Best Practice Checklist", level=1)
    p = doc.add_paragraph()
    run = p.add_run(
        "Use this checklist to manually verify each security control in your F5 "
        "Distributed Cloud environment. Tick the checkbox once you have confirmed "
        "the control is correctly configured. This checklist is generic and applies "
        "to any deployment."
    )
    run.font.size = Pt(10)
    run.font.color.rgb = _GREY

    rules = _dedup_rules(all_findings)

    # Group by category
    by_cat = {}
    for f in rules:
        by_cat.setdefault(_category_label(f.category), []).append(f)

    for cat in sorted(by_cat.keys()):
        items = sorted(
            by_cat[cat],
            key=lambda f: (_SEVERITY_ORDER.get(f.severity.upper(), 9), f.rule_id),
        )
        _heading(doc, cat, level=2, color=_NAVY)

        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _header_row(table, ["Done", "Security Check", "Best Practice / What to Verify", "Reference"])

        for f in items:
            row = table.add_row().cells
            _set_cell_text(row[0], CHECKBOX, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)

            # Check name + rule id + severity
            _set_cell_text(row[1], "", size=9)
            cp = row[1].paragraphs[0]
            r1 = cp.add_run(f.rule_name or f.rule_id)
            r1.bold = True
            r1.font.size = Pt(9)
            meta = cp.add_run(f"\n{f.rule_id} · {f.severity.upper()}")
            meta.font.size = Pt(8)
            meta.font.color.rgb = _GREY

            # Best practice / what to verify
            verify_text = (f.remediation or "").strip()
            if not verify_text:
                verify_text = (f.verify or f.message or "").strip()
            if f.verify and f.remediation:
                verify_text = f"{f.remediation.strip()}\nHow to verify: {f.verify.strip()}"
            _set_cell_text(row[2], verify_text or "Review this control.", size=9)

            # Reference
            _set_cell_text(row[3], (f.reference_url or "").strip(), size=8, color=_GREY)

        _set_col_widths(table, [Inches(0.5), Inches(2.2), Inches(3.4), Inches(1.4)])

    # Manual review items
    if manual_items:
        _heading(doc, "Manual Review Items", level=2, color=_NAVY)
        p = doc.add_paragraph()
        run = p.add_run(
            "These items require a manual conversation or console review — they "
            "cannot be auto-detected from configuration."
        )
        run.font.size = Pt(9)
        run.font.color.rgb = _GREY

        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        _header_row(table, ["Done", "Item", "Recommendation"])
        for item in manual_items:
            row = table.add_row().cells
            _set_cell_text(row[0], CHECKBOX, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
            _set_cell_text(row[1], getattr(item, "label", ""), bold=True, size=9)
            rec = (getattr(item, "recommendation", "") or "").strip()
            ref = getattr(item, "reference_url", "") or ""
            if ref:
                rec = f"{rec}\n{ref}" if rec else ref
            _set_cell_text(row[2], rec, size=9)
        _set_col_widths(table, [Inches(0.5), Inches(3.0), Inches(4.0)])


def _add_appendix(doc, lb_reports, edited_messages=None, excluded_findings=None):
    edited_messages = edited_messages or {}
    excluded_findings = excluded_findings or set()

    doc.add_page_break()
    _heading(doc, "Part 2 — Per-Report Appendix (Your Assessment)", level=1)
    p = doc.add_paragraph()
    run = p.add_run(
        "The tables below list the actual findings from your audit, grouped by "
        "Load Balancer. Use the checkbox to confirm each item has been reviewed "
        "or remediated in your tenant."
    )
    run.font.size = Pt(10)
    run.font.color.rgb = _GREY

    for lb in lb_reports:
        findings = rd.lb_all_findings(lb)
        # Deduplicate by key, drop excluded
        seen = set()
        rows_data = []
        for f in findings:
            key = f"{f.rule_id}::{f.object_name}"
            if key in seen or key in excluded_findings:
                continue
            seen.add(key)
            rows_data.append(f)

        if not rows_data:
            continue

        # Sort: actionable (non-PASS) first, then by severity
        rows_data.sort(
            key=lambda f: (
                0 if f.status.upper() in ("FAIL", "WARN", "ERROR") else 1,
                _SEVERITY_ORDER.get(f.severity.upper(), 9),
                f.rule_id,
            )
        )

        _heading(doc, f"Load Balancer: {lb.name}", level=2, color=_NAVY)

        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        _header_row(table, ["Done", "Status", "Check / Object", "Finding", "Recommendation"])

        for f in rows_data:
            row = table.add_row().cells
            _set_cell_text(row[0], CHECKBOX, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)

            label = rd.status_display_label(f.status)
            _set_cell_text(row[1], label, bold=True, size=9, color=_status_color(f.status, f.severity),
                           align=WD_ALIGN_PARAGRAPH.CENTER)

            _set_cell_text(row[2], "", size=9)
            cp = row[2].paragraphs[0]
            r1 = cp.add_run(f.rule_name or f.rule_id)
            r1.bold = True
            r1.font.size = Pt(9)
            obj = f.object_name if f.object_name and f.object_name != "unknown" else ""
            if obj:
                r2 = cp.add_run(f"\n{obj}")
                r2.font.size = Pt(8)
                r2.font.color.rgb = _GREY

            key = f"{f.rule_id}::{f.object_name}"
            msg = edited_messages.get(key, f.message)
            _set_cell_text(row[3], (msg or "").strip(), size=9)

            rec = (f.remediation or "").strip()
            if f.reference_url:
                rec = f"{rec}\n{f.reference_url}" if rec else f.reference_url
            _set_cell_text(row[4], rec, size=9)

        _set_col_widths(table, [Inches(0.45), Inches(0.7), Inches(1.8), Inches(2.35), Inches(2.2)])


def _add_cover(doc, summary, lb_count):
    for _ in range(3):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Security Best Practice Guide")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = _NAVY

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("F5 Distributed Cloud — Manual Verification Checklist")
    r.font.size = Pt(13)
    r.font.color.rgb = _GREY

    doc.add_paragraph()
    for label, value in [
        ("Tenant", summary.tenant),
        ("Namespace", summary.namespace),
        ("Load Balancers Assessed", str(lb_count)),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]:
        line = doc.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rl = line.add_run(f"{label}:  ")
        rl.bold = True
        rl.font.size = Pt(11)
        rv = line.add_run(str(value))
        rv.font.size = Pt(11)

    for _ in range(2):
        doc.add_paragraph()
    conf = doc.add_paragraph()
    conf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rc = conf.add_run("CONFIDENTIAL")
    rc.font.size = Pt(10)
    rc.font.color.rgb = _RED
    doc.add_page_break()


# ----------------------------------------------------------------------
# Public API
# ----------------------------------------------------------------------

def render_best_practice_docx(summary, lb_reports, all_findings,
                              manual_items=None, edited_messages=None,
                              excluded_findings=None) -> bytes:
    """Build the combined best-practice guide + per-report appendix, return bytes."""
    doc = Document()

    # Base font
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    _add_cover(doc, summary, len(lb_reports))
    _add_generic_guide(doc, all_findings, manual_items)
    _add_appendix(doc, lb_reports, edited_messages=edited_messages,
                  excluded_findings=excluded_findings)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_checklist_cover(doc):
    """Generic, customer-facing cover — no tenant-specific data."""
    for _ in range(3):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Security Review Checklist")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = _NAVY

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("F5 Distributed Cloud — Per Load Balancer Security Controls")
    r.font.size = Pt(13)
    r.font.color.rgb = _GREY

    doc.add_paragraph()
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ri = intro.add_run(
        "This document lists every security control we review on each Load "
        "Balancer. Use it to verify your own configuration and tick each "
        "control once confirmed."
    )
    ri.font.size = Pt(11)
    ri.font.color.rgb = _GREY

    for _ in range(2):
        doc.add_paragraph()
    date_line = doc.add_paragraph()
    date_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rd_ = date_line.add_run("Generated: " + datetime.now().strftime("%Y-%m-%d"))
    rd_.font.size = Pt(10)
    rd_.font.color.rgb = _GREY
    doc.add_page_break()


def render_checklist_docx(all_findings, manual_items=None) -> bytes:
    """Generic customer-facing checklist only — no tenant data, no appendix.

    Lists the full catalog of per-LB security controls with printable
    checkboxes so a customer can self-verify their environment.
    """
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    _add_checklist_cover(doc)
    _add_generic_guide(doc, all_findings, manual_items)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
